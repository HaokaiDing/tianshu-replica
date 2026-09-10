import html
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BODY_FONT = "Noto Sans CJK SC"


def cells(line: str):
    return [item.strip() for item in line.strip().split("|")[1:-1]]


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(inches * 1440)))
    width.set(qn("w:type"), "dxa")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def put_text(cell, text, *, bold=False, size=8, color=None, center=False):
    cell.text = ""
    for index, part in enumerate(text.split("<br>")):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run(part)
        run.bold = bold
        run.font.name = BODY_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)


def parse_markdown(markdown: str):
    lines = [line for line in markdown.splitlines() if line.strip()]
    heading = next((line.lstrip("# ").strip() for line in lines if line.startswith("#")), "分镜剧本")
    header_at = next(index for index, line in enumerate(lines) if line.startswith("|") and "镜头号" in line)
    header = cells(lines[header_at])
    rows = [cells(line) for line in lines[header_at + 2 :] if line.startswith("|")]
    return heading, header, rows


def add_storyboard(doc, heading, header, rows, *, page_break_before=False):
    if len(header) != 7 or any(len(row) != 7 for row in rows):
        raise ValueError("expected a strict 7-column storyboard table")
    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.paragraph_format.page_break_before = page_break_before
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(heading)
    run.bold = True
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(5)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = False
    widths = [0.60, 2.30, 2.25, 1.35, 1.50, 1.95, 0.55]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    header_row = table.rows[0]
    repeat_header(header_row)
    for index, value in enumerate(header):
        cell = header_row.cells[index]
        set_width(cell, widths[index])
        set_cell_margins(cell)
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        put_text(cell, value, bold=True, size=7, color=(255, 255, 255), center=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_width(cell, widths[index])
            set_cell_margins(cell)
            if row_index % 2:
                shade(cell, "F3F6FA")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            put_text(cell, value, size=8, center=index in (0, 6))
def main(input_path: Path, output_path: Path):
    chunks = [chunk for chunk in input_path.read_text(encoding="utf-8").split("\n---\n") if chunk.strip()]
    parsed = [parse_markdown(chunk) for chunk in chunks]
    if any(len(header) != 7 or any(len(row) != 7 for row in rows) for _, header, rows in parsed):
        raise ValueError("expected a strict 7-column storyboard table")
    doc = Document()
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    for border in list(title_properties.findall(qn("w:pBdr"))):
        title_properties.remove(border)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.28)
    section.bottom_margin = Inches(0.28)
    section.left_margin = Inches(0.22)
    section.right_margin = Inches(0.22)
    for index, (heading, header, rows) in enumerate(parsed):
        add_storyboard(doc, heading, header, rows, page_break_before=index > 0)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
